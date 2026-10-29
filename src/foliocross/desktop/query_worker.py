"""Short-lived, read-only desktop queries isolated from the Qt event loop."""

import argparse
import json
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document

from ..application.service import FolioCrossService
from ..errors import FolioCrossError
from ..infrastructure.workspace import Workspace


def preview_artifact(path: str) -> str:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        with source.open(encoding="utf-8-sig") as stream:
            return stream.read(40_000)
    if suffix in {".html", ".htm"}:
        return BeautifulSoup(source.read_bytes(), "lxml").get_text("\n", strip=True)[:40_000]
    if suffix == ".docx":
        document = Document(source)
        paragraphs = list(document.paragraphs)
        paragraphs.extend(
            p
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
        )
        return "\n\n".join(p.text for p in paragraphs if p.text.strip())[:40_000]
    if suffix == ".epub":
        pieces = []
        size = 0
        with zipfile.ZipFile(source) as archive:
            for name in archive.namelist():
                if Path(name).suffix.lower() not in {".xhtml", ".html", ".htm"}:
                    continue
                text = BeautifulSoup(archive.read(name), "lxml").get_text("\n", strip=True)
                pieces.append(text[: 40_000 - size])
                size += len(text) + 2
                if size >= 40_000:
                    break
        return "\n\n".join(pieces)[:40_000]
    raise FolioCrossError("preview_unsupported", "这种文件暂不支持应用内预览。")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", required=True)
    parser.add_argument(
        "--action", choices=["preview", "review_estimate", "affected"], required=True
    )
    parser.add_argument("--arguments", required=True)
    args = parser.parse_args()
    try:
        values = json.loads(args.arguments)
        if args.action == "preview":
            result = {"text": preview_artifact(values["path"])}
        else:
            service = FolioCrossService(Workspace(args.workspace))
            if args.action == "review_estimate":
                result = service.estimate_deep_review(values["job_id"])
            else:
                result = {
                    "rows": service.repository.affected_segments(values["job_id"], values["source"])
                }
        print(json.dumps({"ok": True, **result}, ensure_ascii=True))
    except Exception as error:
        print(json.dumps({"ok": False, "error": {"message": str(error)}}, ensure_ascii=True))


if __name__ == "__main__":
    main()
