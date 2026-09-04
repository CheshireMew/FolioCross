"""Output dispatch for translated jobs."""

from __future__ import annotations

import html
import re
import zipfile
from pathlib import Path

from docx import Document

from ...errors import FolioCrossError
from .docx import export_docx_template
from .epub import export_epub_template
from .htmlish import fill_html_template
from .markers import MARKER_PATTERN

SUPPORTED_OUTPUTS = {"epub", "docx", "html", "txt", "markdown"}


def _chapter_rows(repository, job_id: str) -> list[tuple[dict, list[dict]]]:
    result = []
    for chapter in repository.chapters_for_job(job_id):
        result.append((chapter, repository.segments_for_chapter(chapter["id"])))
    return result


def _display_text(segment: dict, bilingual: bool) -> str:
    target = MARKER_PATTERN.sub("", segment["target_text"] or segment["source_text"])
    source = MARKER_PATTERN.sub("", segment["source_text"])
    return f"{target}\n{source}" if bilingual else target


def _chapter_content(chapter: dict, segments: list[dict]) -> tuple[str, list[dict]]:
    """Use a translated leading heading as the chapter title instead of duplicating it."""

    if not segments:
        return chapter["title"], segments
    first = segments[0]
    source_heading = re.sub(r"^#{1,6}\s+", "", MARKER_PATTERN.sub("", first["source_text"]).strip())
    if source_heading != chapter["title"].strip():
        return chapter["title"], segments
    translated = re.sub(
        r"^#{1,6}\s+",
        "",
        MARKER_PATTERN.sub("", first["target_text"] or source_heading).strip(),
    )
    return translated or chapter["title"], segments[1:]


def _book_display_title(title: str, chapters: list[tuple[dict, list[dict]]]) -> str:
    if not chapters or chapters[0][0]["title"].strip() != title.strip():
        return title
    translated, _ = _chapter_content(*chapters[0])
    return translated


def _write_txt(
    path: str, chapters: list[tuple[dict, list[dict]]], *, markdown: bool, bilingual: bool
) -> str:
    pieces: list[str] = []
    for chapter, segments in chapters:
        title, content = _chapter_content(chapter, segments)
        pieces.append(f"# {title}" if markdown else title)
        if bilingual and title != chapter["title"]:
            pieces.append(chapter["title"])
        pieces.extend(_display_text(segment, bilingual) for segment in content)
    Path(path).write_text("\n\n".join(pieces) + "\n", encoding="utf-8")
    return path


def _write_html(
    path: str, title: str, chapters: list[tuple[dict, list[dict]]], *, bilingual: bool
) -> str:
    body = []
    for chapter, segments in chapters:
        chapter_title, content = _chapter_content(chapter, segments)
        body.append(f"<section><h2>{html.escape(chapter_title)}</h2>")
        if bilingual and chapter_title != chapter["title"]:
            body.append(f'<p class="source">{html.escape(chapter["title"])}</p>')
        for segment in content:
            target = html.escape(
                MARKER_PATTERN.sub("", segment["target_text"] or segment["source_text"])
            )
            body.append(f'<p class="target">{target}</p>')
            if bilingual:
                source = html.escape(MARKER_PATTERN.sub("", segment["source_text"]))
                body.append(f'<p class="source">{source}</p>')
        body.append("</section>")
    document = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>body{{max-width:48rem;margin:3rem auto;padding:0 1.5rem;font:18px/1.8 serif;color:#241f1a}}
.source{{color:#777;font-size:.9em;border-left:3px solid #ddd;padding-left:1em}} h2{{margin-top:2.5em}}</style>
</head><body>{"".join(body)}</body></html>"""
    Path(path).write_text(document, encoding="utf-8")
    return path


def _write_docx(
    path: str, title: str, chapters: list[tuple[dict, list[dict]]], *, bilingual: bool
) -> str:
    document = Document()
    document.core_properties.title = title
    document.add_heading(title, level=0)
    for chapter, segments in chapters:
        chapter_title, content = _chapter_content(chapter, segments)
        document.add_heading(chapter_title, level=1)
        if bilingual and chapter_title != chapter["title"]:
            document.add_paragraph(chapter["title"], style="Subtitle")
        for segment in content:
            document.add_paragraph(_display_text(segment, bilingual))
    document.save(path)
    return path


def _write_generated_epub(
    path: str,
    title: str,
    author: str,
    chapters: list[tuple[dict, list[dict]]],
    *,
    bilingual: bool,
) -> str:
    container = """<?xml version="1.0"?><container version="1.0"
 xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>
 <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
</rootfiles></container>"""
    manifest = []
    spine = []
    nav = []
    chapter_files: dict[str, str] = {}
    for index, (chapter, segments) in enumerate(chapters):
        chapter_title, content = _chapter_content(chapter, segments)
        name = f"chapter-{index + 1}.xhtml"
        manifest.append(f'<item id="c{index}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{index}"/>')
        nav.append(f'<li><a href="{name}">{html.escape(chapter_title)}</a></li>')
        paragraphs = "".join(
            f"<p>{html.escape(_display_text(segment, bilingual))}</p>" for segment in content
        )
        source_title = (
            f'<p class="source">{html.escape(chapter["title"])}</p>'
            if bilingual and chapter_title != chapter["title"]
            else ""
        )
        chapter_files[name] = (
            '<?xml version="1.0" encoding="utf-8"?>'
            '<html xmlns="http://www.w3.org/1999/xhtml" lang="zh-CN"><head>'
            f"<title>{html.escape(chapter_title)}</title></head><body>"
            f"<h1>{html.escape(chapter_title)}</h1>{source_title}{paragraphs}</body></html>"
        )
    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">foliocross</dc:identifier>
<dc:title>{html.escape(title)}</dc:title><dc:creator>{html.escape(author)}</dc:creator><dc:language>zh-CN</dc:language>
<meta property="dcterms:modified">2026-09-03T00:00:00Z</meta></metadata>
<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>{"".join(manifest)}</manifest>
<spine>{"".join(spine)}</spine></package>"""
    nav_xhtml = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
        f"<head><title>{html.escape(title)}</title></head><body><nav epub:type="
        f'"toc"><ol>{"".join(nav)}</ol></nav></body></html>'
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", opf)
        archive.writestr("OEBPS/nav.xhtml", nav_xhtml)
        for name, content in chapter_files.items():
            archive.writestr(f"OEBPS/{name}", content)
    return path


def export_job(
    repository, job_id: str, format_name: str | None = None, output_path: str | None = None
) -> str:
    job = repository.get_job(job_id)
    config = job["config"]
    format_name = (format_name or config.get("output_format") or "epub").lower()
    if format_name == "md":
        format_name = "markdown"
    if format_name not in SUPPORTED_OUTPUTS:
        raise FolioCrossError("unsupported_output", f"不支持输出格式：{format_name}")
    chapters = _chapter_rows(repository, job_id)
    missing = sum(1 for _, segments in chapters for item in segments if not item["target_text"])
    if missing:
        raise FolioCrossError(
            "translation_incomplete", f"还有 {missing} 个段落尚未翻译，不能生成最终成品。"
        )
    book = repository.get_book(job["book_id"])
    suffix = ".md" if format_name == "markdown" else f".{format_name}"
    if output_path is None:
        output_dir = repository.workspace.book_export_dir(book["id"])
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = str(output_dir / f"{_safe_name(book['title'])}.zh-CN{suffix}")
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    bilingual = bool(config.get("bilingual", False))
    display_title = _book_display_title(book["title"], chapters)
    source_format = book["source_format"]
    template = book["working_source_path"]
    all_segments = [item for _, segments in chapters for item in segments]
    if format_name == "epub" and source_format in {"epub", "mobi", "azw", "azw3", "azw4", "prc"}:
        return export_epub_template(
            template, output_path, all_segments, target_lang="zh-CN", bilingual=bilingual
        )
    if format_name == "docx" and source_format == "docx":
        return export_docx_template(template, output_path, all_segments, bilingual=bilingual)
    if format_name == "html" and source_format == "html":
        return fill_html_template(template, output_path, all_segments, bilingual=bilingual)
    if format_name == "txt":
        return _write_txt(output_path, chapters, markdown=False, bilingual=bilingual)
    if format_name == "markdown":
        return _write_txt(output_path, chapters, markdown=True, bilingual=bilingual)
    if format_name == "html":
        return _write_html(output_path, display_title, chapters, bilingual=bilingual)
    if format_name == "docx":
        return _write_docx(output_path, display_title, chapters, bilingual=bilingual)
    if format_name == "epub":
        return _write_generated_epub(
            output_path, display_title, book["author"], chapters, bilingual=bilingual
        )
    raise AssertionError(format_name)


def _safe_name(value: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if character in invalid else character for character in value).strip(" .")
    return cleaned or "book"
