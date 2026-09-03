"""DOCX parser and conservative template fillback."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Iterable

from docx import Document

from ...domain.models import ChapterDraft, ParsedBook, SegmentDraft


def _all_paragraphs(document: Document):
    for index, paragraph in enumerate(document.paragraphs):
        yield paragraph, {"kind": "paragraph", "paragraph": index}
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield (
                        paragraph,
                        {
                            "kind": "table_cell",
                            "table": table_index,
                            "row": row_index,
                            "cell": cell_index,
                            "paragraph": paragraph_index,
                        },
                    )


def parse_docx(path: str, source_lang: str = "auto") -> ParsedBook:
    document = Document(path)
    title = document.core_properties.title or Path(path).stem
    author = document.core_properties.author or ""
    chapters: list[ChapterDraft] = []
    current_title = title
    current: list[SegmentDraft] = []
    warnings: list[str] = []
    for paragraph, locator in _all_paragraphs(document):
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style.lower().startswith("heading") and current:
            chapters.append(ChapterDraft(current_title, current, {"source": "docx"}))
            current = []
        if style.lower().startswith("heading"):
            current_title = text
        locator["style"] = style
        current.append(SegmentDraft(text, locator))
    if current:
        chapters.append(ChapterDraft(current_title, current, {"source": "docx"}))
    if not chapters:
        raise ValueError("DOCX 中没有可翻译正文")
    if document.element.xpath(".//*[local-name()='txbxContent']"):
        warnings.append("检测到文本框；首版不会翻译文本框内文字。")
    if document.element.xpath(".//*[local-name()='ins' or local-name()='del']"):
        warnings.append("检测到修订记录；首版只翻译接受后可见的正文。")
    return ParsedBook(title, author, "docx", source_lang, chapters, warnings=warnings)


def _paragraph_for_locator(document: Document, locator: dict):
    if locator.get("kind") == "paragraph":
        return document.paragraphs[int(locator["paragraph"])]
    table = document.tables[int(locator["table"])]
    cell = table.rows[int(locator["row"])].cells[int(locator["cell"])]
    return cell.paragraphs[int(locator["paragraph"])]


def _replace_paragraph(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    runs = list(paragraph.runs)
    source_lengths = [max(1, len(run.text)) for run in runs]
    total = sum(source_lengths)
    cursor = 0
    for index, run in enumerate(runs):
        if index == len(runs) - 1:
            run.text = text[cursor:]
        else:
            end = round(len(text) * sum(source_lengths[: index + 1]) / total)
            run.text = text[cursor:end]
            cursor = end


def export_docx_template(
    source_path: str, output_path: str, segments: Iterable[dict], *, bilingual: bool = False
) -> str:
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)
    document = Document(output_path)
    for segment in segments:
        target = segment["target_text"] or segment["source_text"]
        if bilingual:
            target = f"{target}\n{segment['source_text']}"
        _replace_paragraph(_paragraph_for_locator(document, segment["locator"]), target)
    document.save(output_path)
    return output_path
