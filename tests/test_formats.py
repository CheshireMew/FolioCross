from __future__ import annotations

import base64
import hashlib
import zipfile
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from conftest import write_sample_epub
from docx import Document

from foliocross.errors import FolioCrossError
from foliocross.infrastructure.formats.common import split_long_segments
from foliocross.infrastructure.formats.docx import export_docx_template, parse_docx
from foliocross.infrastructure.formats.epub import export_epub_template, parse_epub
from foliocross.infrastructure.formats.fb2 import parse_fb2
from foliocross.infrastructure.formats.htmlish import content_blocks, fill_html_template, parse_html
from foliocross.infrastructure.formats.markers import (
    render_fragment,
    repair_missing_markers,
    restore_protected_markers,
    serialize_block,
    strip_protected_markers,
    validate_markers,
)
from foliocross.infrastructure.formats.text import parse_text

ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def test_epub_parser_preserves_inline_and_note_markers(tmp_path: Path):
    source = write_sample_epub(tmp_path / "book.epub")
    book = parse_epub(str(source))
    assert book.title == "Sample Book"
    assert book.author == "Tester"
    assert book.source_lang == "en"
    marker_segments = [
        segment for chapter in book.chapters for segment in chapter.segments if segment.markers
    ]
    assert marker_segments
    kinds = {marker["kind"] for segment in marker_segments for marker in segment.markers}
    assert {"annotation", "tag_start", "tag_end", "inline_atom"} <= kinds


def test_epub_template_export_keeps_resources_and_restores_markup(tmp_path: Path):
    source = write_sample_epub(tmp_path / "book.epub")
    book = parse_epub(str(source))
    segments = []
    for chapter in book.chapters:
        for index, draft in enumerate(chapter.segments):
            segments.append(
                {
                    "source_text": draft.source_text,
                    "target_text": "译文" + draft.source_text,
                    "locator": draft.locator,
                    "markers": draft.markers,
                    "segment_index": index,
                }
            )
    output = tmp_path / "translated.epub"
    export_epub_template(str(source), str(output), segments)
    with zipfile.ZipFile(output) as archive:
        assert "OEBPS/style.css" in archive.namelist()
        assert "OEBPS/cover.png" in archive.namelist()
        content = archive.read("OEBPS/chapter.xhtml").decode("utf-8")
        assert "⟦FC:" not in content
        assert 'href="#n1"' in content
        assert "<em>important</em>" in content
        opf = archive.read("OEBPS/content.opf").decode("utf-8")
        nav = archive.read("OEBPS/nav.xhtml").decode("utf-8")
        assert "zh-CN" in opf
        assert ">译文Chapter One</a>" in nav


def test_epub_translates_linked_note_resource_outside_spine(tmp_path: Path):
    source = tmp_path / "separate-notes.epub"
    container = """<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
    <rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>"""
    opf = """<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
    <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Notes</dc:title>
    <dc:language>en</dc:language></metadata><manifest>
    <item id="body" href="body.xhtml" media-type="application/xhtml+xml"/>
    <item id="notes" href="notes.xhtml" media-type="application/xhtml+xml"/>
    </manifest><spine><itemref idref="body"/></spine></package>"""
    body = """<html xmlns="http://www.w3.org/1999/xhtml"
    xmlns:epub="http://www.idpf.org/2007/ops"><body><p>Body
    <a epub:type="noteref" href="notes.xhtml#note-1">1</a>.</p></body></html>"""
    notes = """<html xmlns="http://www.w3.org/1999/xhtml"
    xmlns:epub="http://www.idpf.org/2007/ops"><body><aside id="note-1"
    epub:type="footnote"><p>Separate note text.</p></aside></body></html>"""
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", opf)
        archive.writestr("OEBPS/body.xhtml", body)
        archive.writestr("OEBPS/notes.xhtml", notes)

    book = parse_epub(str(source))
    assert len(book.chapters) == 2
    segments = [
        {
            "source_text": segment.source_text,
            "target_text": "译文" + segment.source_text,
            "locator": segment.locator,
            "markers": segment.markers,
        }
        for chapter in book.chapters
        for segment in chapter.segments
    ]
    output = tmp_path / "translated-notes.epub"
    export_epub_template(str(source), str(output), segments)
    with zipfile.ZipFile(output) as archive:
        assert "译文Separate note text." in archive.read("OEBPS/notes.xhtml").decode("utf-8")
        assert 'href="notes.xhtml#note-1"' in archive.read("OEBPS/body.xhtml").decode("utf-8")


def test_marker_validator_detects_loss(tmp_path: Path):
    source = write_sample_epub(tmp_path / "book.epub")
    book = parse_epub(str(source))
    segment = next(s for c in book.chapters for s in c.segments if s.markers)
    assert validate_markers(segment.source_text, segment.markers) == []
    missing = segment.source_text.replace(segment.markers[0]["token"], "")
    assert validate_markers(missing, segment.markers)


def test_manual_visible_edit_restores_protected_markers(tmp_path: Path):
    source = write_sample_epub(tmp_path / "book.epub")
    book = parse_epub(str(source))
    segment = next(s for c in book.chapters for s in c.segments if s.markers)
    visible_edit = "这是一段经过人工改写的重要文字。"

    restored = restore_protected_markers(
        segment.source_text,
        segment.source_text,
        visible_edit + segment.markers[0]["token"],
        segment.markers,
    )

    assert strip_protected_markers(restored) == visible_edit
    assert validate_markers(restored, segment.markers) == []
    assert restored.count(segment.markers[0]["token"]) == 1


def test_nested_inline_annotation_marker_order_is_valid():
    soup = BeautifulSoup(
        '<p>Before <em>important<a epub:type="noteref" href="#n1">1</a></em> after.</p>',
        "html.parser",
    )
    text, markers = serialize_block(soup.p, identity="nested")
    assert validate_markers(text, markers) == []
    assert [marker["kind"] for marker in markers] == [
        "tag_start",
        "annotation",
        "tag_end",
    ]


def test_manual_edit_keeps_markup_attached_to_unchanged_words():
    soup = BeautifulSoup(
        '<p>Before <em>important</em> text<a epub:type="noteref" href="#n1">1</a>.</p>',
        "html.parser",
    )
    source, markers = serialize_block(soup.p, identity="edit-positions")
    edited = "A much longer introduction before important text."
    restored = restore_protected_markers(source, source, edited, markers)
    output = BeautifulSoup(render_fragment(restored, markers), "html.parser")
    assert output.em.get_text() == "important"
    assert str(output).endswith('text<a epub:type="noteref" href="#n1">1</a>.')
    assert strip_protected_markers(restored) == edited
    assert validate_markers(restored, markers) == []


def test_marker_fallback_restores_nested_order_without_changing_visible_text():
    soup = BeautifulSoup("<p>Before <em><strong>important</strong></em> after.</p>", "html.parser")
    source, markers = serialize_block(soup.p, identity="fallback-order")
    restored = repair_missing_markers(source, "之前重要之后。", markers)
    assert validate_markers(restored, markers) == []
    assert strip_protected_markers(restored) == "之前重要之后。"


def test_long_marked_epub_segment_splits_and_reassembles_exactly(tmp_path: Path):
    source = write_sample_epub(tmp_path / "long-marked.epub")
    original = parse_epub(str(source))
    split = split_long_segments(original, 45)
    original_segments = [segment for chapter in original.chapters for segment in chapter.segments]
    split_segments = [segment for chapter in split.chapters for segment in chapter.segments]
    for original_segment in original_segments:
        matching = [
            segment
            for segment in split_segments
            if segment.locator.get("ordinal") == original_segment.locator.get("ordinal")
        ]
        assert "".join(segment.source_text for segment in matching) == original_segment.source_text
        assert [marker["token"] for segment in matching for marker in segment.markers] == [
            marker["token"] for marker in original_segment.markers
        ]

    output = tmp_path / "reassembled.epub"
    export_epub_template(
        str(source),
        str(output),
        [
            {
                "source_text": segment.source_text,
                "target_text": segment.source_text,
                "locator": segment.locator,
                "markers": segment.markers,
            }
            for segment in split_segments
        ],
    )
    assert output.is_file()


def test_text_parser_splits_headings(tmp_path: Path):
    path = tmp_path / "book.md"
    path.write_text("# One\n\nFirst.\n\n# Two\n\nSecond.", encoding="utf-8")
    book = parse_text(str(path))
    assert [chapter.title for chapter in book.chapters] == ["One", "Two"]


def test_html_and_fb2_are_mapped_to_chapters_and_segments(tmp_path: Path):
    html_path = tmp_path / "book.html"
    html_path.write_text(
        "<html><head><title>HTML Book</title></head><body><h1>One</h1><p>Paragraph.</p></body></html>",
        encoding="utf-8",
    )
    html_book = parse_html(str(html_path))
    assert html_book.title == "HTML Book"
    assert any(segment.source_text == "Paragraph." for segment in html_book.chapters[0].segments)

    fb2_path = tmp_path / "book.fb2"
    fb2_path.write_text(
        """<?xml version="1.0" encoding="utf-8"?>
        <FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">
        <description><title-info><book-title>FB2 Book</book-title><author><first-name>A</first-name>
        <last-name>Writer</last-name></author></title-info></description>
        <body><section><title><p>Chapter</p></title><p>First text.</p><p>Second text.</p></section></body>
        </FictionBook>""",
        encoding="utf-8",
    )
    fb2_book = parse_fb2(str(fb2_path))
    assert fb2_book.title == "FB2 Book"
    assert fb2_book.chapters[0].title == "Chapter"
    assert len(fb2_book.chapters[0].segments) == 2


def test_nested_html_blocks_are_not_extracted_twice(tmp_path: Path):
    html_path = tmp_path / "nested.html"
    html_path.write_text(
        "<html><body><blockquote><p>Quoted once.</p></blockquote>"
        "<ul><li><p>Listed once.</p></li></ul>"
        "<p><img src='cover.png'/></p></body></html>",
        encoding="utf-8",
    )
    book = parse_html(str(html_path))
    texts = [segment.source_text for segment in book.chapters[0].segments]
    assert sum("Quoted once." in text for text in texts) == 1
    assert sum("Listed once." in text for text in texts) == 1
    assert len(texts) == 2

    soup = BeautifulSoup(html_path.read_bytes(), "lxml")
    assert [block.name for block in content_blocks(soup)] == ["p", "p", "p"]


def test_parent_list_text_and_nested_list_are_both_preserved_once(tmp_path: Path):
    html_path = tmp_path / "nested-list.html"
    html_path.write_text(
        "<html><body><ul><li>Parent text<ul><li>Child text</li></ul></li></ul></body></html>",
        encoding="utf-8",
    )
    book = parse_html(str(html_path))
    texts = [segment.source_text for segment in book.chapters[0].segments]
    assert texts == ["Parent text", "Child text"]


def test_nested_fb2_sections_are_not_duplicated(tmp_path: Path):
    path = tmp_path / "nested.fb2"
    path.write_text(
        """<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0"><body>
        <section><title><p>Outer</p></title><p>Outer text.</p>
        <section><title><p>Inner</p></title><p>Inner text.</p></section>
        </section></body></FictionBook>""",
        encoding="utf-8",
    )
    book = parse_fb2(str(path))
    assert [chapter.title for chapter in book.chapters] == ["Outer", "Inner"]
    texts = [segment.source_text for chapter in book.chapters for segment in chapter.segments]
    assert texts == ["Outer text.", "Inner text."]


def test_template_export_rejects_missing_or_stale_mapping(tmp_path: Path):
    epub = write_sample_epub(tmp_path / "book.epub")
    with pytest.raises(FolioCrossError, match="EPUB 原位定位"):
        export_epub_template(
            str(epub),
            str(tmp_path / "bad.epub"),
            [{"source_text": "x", "target_text": "译", "locator": {}, "markers": []}],
        )

    html = tmp_path / "book.html"
    html.write_text("<html><body><p>Original.</p></body></html>", encoding="utf-8")
    with pytest.raises(FolioCrossError, match="原文不一致"):
        fill_html_template(
            str(html),
            str(tmp_path / "bad.html"),
            [
                {
                    "source_text": "Changed.",
                    "target_text": "译文。",
                    "locator": {"ordinal": 0},
                    "markers": [],
                }
            ],
        )


def test_docx_template_keeps_structure_and_replaces_text(tmp_path: Path):
    source = tmp_path / "book.docx"
    document = Document()
    document.add_heading("Heading", 1)
    paragraph = document.add_paragraph()
    bold = paragraph.add_run("Bold")
    bold.bold = True
    paragraph.add_run(" normal")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell text"
    document.save(source)
    book = parse_docx(str(source))
    drafts = [segment for chapter in book.chapters for segment in chapter.segments]
    rows = [
        {"source_text": item.source_text, "target_text": f"译文{index}", "locator": item.locator}
        for index, item in enumerate(drafts)
    ]
    output = tmp_path / "translated.docx"
    export_docx_template(str(source), str(output), rows)
    translated = Document(output)
    assert translated.paragraphs[0].text == "译文0"
    assert translated.paragraphs[1].text == "译文1"
    assert translated.paragraphs[1].runs[0].bold is True
    assert translated.tables[0].cell(0, 0).text == "译文2"

    with pytest.raises(FolioCrossError, match="DOCX 原位定位"):
        export_docx_template(
            str(source),
            str(tmp_path / "bad.docx"),
            [{"source_text": "Heading", "target_text": "标题", "locator": {}}],
        )


def test_docx_export_preserves_package_images_lists_tables_and_unknown_parts(tmp_path: Path):
    source = tmp_path / "rich.docx"
    image_path = tmp_path / "pixel.png"
    image_path.write_bytes(ONE_PIXEL_PNG)
    document = Document()
    document.core_properties.title = "Rich Book"
    document.core_properties.author = "Writer"
    document.add_heading("Chapter", 1)
    document.add_paragraph("First item", style="List Number")
    pictured = document.add_paragraph("Illustrated text")
    pictured.add_run().add_picture(str(image_path))
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table text"
    document.save(source)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("customXml/foliocross-test.bin", b"opaque-extension-data")

    with zipfile.ZipFile(source) as archive:
        media_name = next(name for name in archive.namelist() if name.startswith("word/media/"))
        media_digest = hashlib.sha256(archive.read(media_name)).hexdigest()
        numbering = archive.read("word/numbering.xml")

    parsed = parse_docx(str(source))
    assert parsed.title == "Rich Book"
    assert parsed.author == "Writer"
    drafts = [segment for chapter in parsed.chapters for segment in chapter.segments]
    output = tmp_path / "rich-translated.docx"
    export_docx_template(
        str(source),
        str(output),
        [
            {
                "source_text": item.source_text,
                "target_text": f"译文 {index}",
                "locator": item.locator,
            }
            for index, item in enumerate(drafts)
        ],
    )

    with zipfile.ZipFile(output) as archive:
        assert archive.read("customXml/foliocross-test.bin") == b"opaque-extension-data"
        assert hashlib.sha256(archive.read(media_name)).hexdigest() == media_digest
        assert archive.read("word/numbering.xml") == numbering
    translated = Document(output)
    assert translated.paragraphs[1].style.name == "List Number"
    assert translated.tables[0].cell(0, 0).text.startswith("译文")
    assert translated.inline_shapes
